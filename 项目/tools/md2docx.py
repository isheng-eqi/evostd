"""Markdown → Word(.docx)（零第三方依赖之外的 python-docx）。

只覆盖本项目文档实际用到的语法：
  # / ## / ### 标题、段落、**加粗**、`行内代码`、> 引用、--- 分隔、
  管道表格、``` 代码块、![alt](相对路径) 图片。

用法：
    python tools/md2docx.py 项目/submission/开发设计文档-六节版.md 项目/submission/开发设计文档-六节版.docx
"""
from __future__ import annotations

import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CJK_FONT = "微软雅黑"
MONO_FONT = "Consolas"


def _set_cjk(run, font=CJK_FONT):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)


def _shade(paragraph, hex_fill="F6F8FA"):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    ppr.append(shd)


def _add_rich_text(paragraph, text: str, size=10.5):
    """处理 **加粗** 与 `行内代码`。"""
    tokens = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            run = paragraph.add_run(tok[2:-2])
            run.bold = True
            run.font.size = Pt(size)
            _set_cjk(run)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            run = paragraph.add_run(tok[1:-1])
            run.font.name = MONO_FONT
            run.font.size = Pt(size - 1)
            _set_cjk(run, MONO_FONT)
        else:
            run = paragraph.add_run(tok)
            run.font.size = Pt(size)
            _set_cjk(run)


def _flush_paragraph(doc, buf: list[str]):
    text = " ".join(x.strip() for x in buf if x.strip())
    if not text:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    _add_rich_text(p, text)


def _add_table(doc, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            val = row[j] if j < len(row) else ""
            _add_rich_text(p, val, size=9.5)
            if i == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def convert(md_path: str, out_path: str):
    base = os.path.dirname(os.path.abspath(md_path))
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = CJK_FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)

    buf: list[str] = []
    table_buf: list[list[str]] = []
    in_code = False

    def flush_table():
        nonlocal table_buf
        if table_buf:
            _add_table(doc, table_buf)
            table_buf = []

    for raw in lines:
        line = raw.rstrip()

        if line.startswith("```"):
            if in_code:
                in_code = False
            else:
                _flush_paragraph(doc, buf)
                buf = []
                in_code = True
            continue

        if in_code:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(line if line else " ")
            run.font.name = MONO_FONT
            run.font.size = Pt(8.5)
            _set_cjk(run, MONO_FONT)
            _shade(p)
            continue

        if line.startswith("|") and line.endswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue
            table_buf.append(cells)
            continue
        else:
            flush_table()

        if re.match(r"^#{1,4}\s", line):
            _flush_paragraph(doc, buf)
            buf = []
            level = len(line) - len(line.lstrip("#"))
            text = line[level:].strip()
            h = doc.add_heading(level=min(level, 3))
            run = h.add_run(text)
            _set_cjk(run)
            run.font.size = Pt(18 - 2 * level)
            run.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)
            continue

        if line.strip() == "---":
            _flush_paragraph(doc, buf)
            buf = []
            continue

        m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line.strip())
        if m:
            _flush_paragraph(doc, buf)
            buf = []
            alt, rel = m.group(1), m.group(2)
            img = rel if os.path.isabs(rel) else os.path.join(base, rel)
            if os.path.exists(img):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(img, width=Cm(15.5))
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(alt)
                r.italic = True
                r.font.size = Pt(9)
                _set_cjk(r)
            else:
                p = doc.add_paragraph(f"[缺图] {rel}")
                p.runs[0].font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            continue

        if line.startswith(">"):
            _flush_paragraph(doc, buf)
            buf = []
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            r = p.add_run(line.lstrip("> ").strip())
            r.italic = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x57, 0x60, 0x6A)
            _set_cjk(r)
            continue

        if not line.strip():
            _flush_paragraph(doc, buf)
            buf = []
            continue

        buf.append(line)

    _flush_paragraph(doc, buf)
    flush_table()
    doc.save(out_path)
    print(f"[md2docx] {md_path} -> {out_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("用法: python tools/md2docx.py <in.md> <out.docx>")
        raise SystemExit(2)
    convert(sys.argv[1], sys.argv[2])
