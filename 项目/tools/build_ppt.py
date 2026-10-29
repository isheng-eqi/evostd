"""生成演示 PPT（16:9，12 页）+ 每页讲稿（写入备注）。

设计原则
--------
- 幻灯片只放**关键词和证据**，话由人讲；
- 每页备注里是一段**口语讲稿**（不是书面语），录视频时照着说即可；
- 图片按原比例缩放，不拉伸。

用法：
    cd 项目
    python tools/build_ppt.py
输出：
    submission/演示PPT.pptx
"""
from __future__ import annotations

import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SUB = os.path.join(ROOT, "submission")
OUT = os.path.join(SUB, "演示PPT.pptx")

W, H = Inches(13.333), Inches(7.5)
INK = RGBColor(0x11, 0x18, 0x27)
MUTED = RGBColor(0x6B, 0x72, 0x80)
ACCENT = RGBColor(0x1F, 0x6F, 0xEB)
RED = RGBColor(0xCF, 0x22, 0x2E)
GREEN = RGBColor(0x1A, 0x7F, 0x37)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x0B, 0x12, 0x20)
SOFT = RGBColor(0xF6, 0xF8, 0xFA)

FONT = "微软雅黑"


def set_font(run, size, color, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if ea is None:
        from lxml import etree
        ea = etree.SubElement(rpr, "{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    ea.set("typeface", FONT)


def textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines: [(text, size, color, bold), ...]"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, (text, size, color, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(8)
        r = p.add_run()
        r.text = text
        set_font(r, size, color, bold)
    return tb


def rect(slide, x, y, w, h, fill, line=None, radius=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    if radius:
        try:
            shape.adjustments[0] = 0.08
        except Exception:
            pass
    return shape


def image_fit(slide, path, x, y, w, h):
    if not os.path.exists(path):
        return False
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    nw, nh = int(iw * scale), int(ih * scale)
    nx, ny = x + (w - nw) // 2, y + (h - nh) // 2
    slide.shapes.add_picture(path, nx, ny, nw, nh)
    # 细边框
    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, nx, ny, nw, nh)
    frame.fill.background()
    frame.line.color.rgb = RGBColor(0xD0, 0xD7, 0xDE)
    frame.line.width = Pt(0.75)
    frame.shadow.inherit = False
    return True


def image_cropped(slide, rel_path, x, y, w, h, crop=None):
    """按比例裁掉截图里无关的区域（侧栏/历史列表），只留要看的部分。"""
    src = os.path.join(SUB, rel_path)
    if not os.path.exists(src):
        return False
    if crop:
        with Image.open(src) as im:
            iw, ih = im.size
            box = (int(iw * crop[0]), int(ih * crop[1]), int(iw * crop[2]), int(ih * crop[3]))
            out_dir = os.path.join(SUB, "screenshots", "_ppt")
            os.makedirs(out_dir, exist_ok=True)
            out = os.path.join(out_dir, os.path.basename(rel_path).replace(".png", "-crop.png"))
            im.crop(box).save(out)
            src = out
    return image_fit(slide, src, x, y, w, h)


def bg(slide, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def page_no(slide, n, total, dark=False):
    textbox(slide, W - Inches(1.1), H - Inches(0.55), Inches(0.6), Inches(0.3),
            [(f"{n} / {total}", 10, RGBColor(0x9C, 0xA3, 0xAF) if dark else MUTED, False)],
            align=PP_ALIGN.RIGHT)


def notes(slide, script, action=""):
    text = script if not action else script + "\n\n【动作】" + action
    slide.notes_slide.notes_text_frame.text = text


SLIDES = [
    dict(kind="cover",
         title="标准驱动的\n可进化决策智能体",
         kicker="第十一届华为 ICT 大赛创新赛 · 赛题 3",
         sub="WarrantedBelief · 太原理工大学",
         script="大家好，我们是 WarrantedBelief 队。我们做的是赛题三，一个可进化的决策智能体。"
                "先说清楚它不是什么——它不是让模型自己给自己打分。它判断对错的依据，来自行业标准的原文。",
         action="停一秒，再点下一页。"),
    dict(kind="statement",
         kicker="问题",
         title="沉睡的数据资产，缺的不是存储",
         big="是「什么算对」没有落到机器能执行的地方",
         chips=["设备台账", "点检表", "图纸影像", "规程文本"],
         script="工厂里、政府部门里，躺着很多这样的东西：设备台账、点检表、图纸、规程。"
                "你说缺数据吗？不缺。缺的是「什么算对」。同一条台账，张三说没问题，李四说不行，谁也拿不出依据。"
                "所以这些数据一直不敢用。",
         action="说到四个标签时，手依次点一下。"),
    dict(kind="statement",
         kicker="主张",
         title="自进化 = 系统扩大它能够分辨的对错范围",
         big="最终裁决者，永远在系统之外",
         bullets=["硬锚点：从条款原文机械抽出的「确定对 / 确定错」样本",
                  "新判据必须先在锚点上证明自己：不误报、能抓到东西",
                  "判据库、锚点、台账都在 MCP 里：模型只能调用，改不了"],
         script="我们的做法，说人话就是：把标准里的条款，变成机器能执行的判据。"
                "比如「检修周期不能超过九十天」，就变成一条能跑的检查。"
                "然后关键的一步——我们从条款原文里，机械地抽出「什么肯定合规、什么肯定违规」的样本，我们叫它硬锚点。"
                "新规则想进库，必须先在锚点上证明自己：不误报、能抓到东西。最后拍板的，是人，不是模型。",
         action="说到「硬锚点」时，点一下第二行。"),
    dict(kind="diagram",
         kicker="系统全貌",
         title="三层分工：模型编排、判定在 MCP、检索在知识库",
         bullets=["上层：Nexent 智能体（deepseek-flash）",
                  "中层：自研 MCP 服务，36 个工具",
                  "右侧：知识库 mfg_assets（10 文档 / 51 分块）",
                  "边界：模型只能调用，改不了判据库"],
         script="整个系统分三层。上面是 Nexent 里的智能体，它负责选工具、串流程、把结果讲成人话。"
                "中间是自研的 MCP 服务，三十六个工具，判据库、锚点、台账都在这里面。右边是知识库，存的是语料。"
                "注意这个边界：模型只能调用，改不了判据库。",
         action="用激光笔沿三层从上往下划。"),
    dict(kind="image_right",
         kicker="平台配置",
         title="这不是空壳：每个工具都能点开调",
         image="screenshots/07-智能体工具绑定.png",
         crop=(0.40, 0.10, 0.99, 0.97),
         bullets=["模型：deepseek-flash（开启 thinking）",
                  "36 个自研 MCP 工具 + 1 个知识库检索工具",
                  "5 个技能（校验 / 提案 / 进化 / 溯源 / 登记）",
                  "能力指纹 9daebab9b1e9"],
         script="这是它在平台上的样子。模型是 deepseek-flash，挂了三十六个 MCP 工具、一个知识库检索工具、五个技能。"
                "配置页里能看到每个工具。这些不是演示用的空壳，每一个都能点开调。",
         action="鼠标在工具列表上滚一下即可，别停留太久。"),
    dict(kind="image_right",
         kicker="检索-推理双驱动",
         title="检索找证据，判据判对错",
         image="screenshots/16-检索推理双驱动.png",
         crop=(0.37, 0.04, 1.0, 1.0),
         bullets=["先调 knowledge_base_search 取语料",
                  "再查判据第 45 条（含 limits 表）",
                  "回答同时给：判据原文 + 语料出处（得分 0.710）",
                  "冲突时以判据为准，写进【未决】"],
         script="这一页是「检索-推理双驱动」。问它 A 类设备的检修周期上限是多少，它会先去知识库检索，拿到语料里的证据，"
                "同时去判据库查第四十五条。回答里两条依据都给出来：一条是判据，一条是语料，还带检索得分。"
                "检索负责找证据，判据负责判对错——这两个不能混。",
         action="说到「两条依据」时，指一下截图里的两段。"),
    dict(kind="image_right",
         kicker="端到端决策",
         title="一条资产，四段式决策记录",
         image="screenshots/11-决策记录-结论与依据.png",
         crop=(0.37, 0.04, 1.0, 1.0),
         bullets=["结论：不能入库",
                  "依据：第 45 条，A 类上限 90 天，实际 365 天",
                  "证据字段原样给出：{field, value, key, max}",
                  "下一步：修正后重新校验"],
         script="再看一条真实的决策。问它这条资产能不能入库。它没有直接回答，而是依次调工具：取资产、逐条校验、尝试入库、查待裁决。"
                "结论是不能入库，依据是第四十五条：A 类设备上限九十天，这条填了三百六十五天。证据字段原样贴出来。",
         action="说到「三百六十五天」时，停半秒。"),
    dict(kind="image_right",
         kicker="弃权可见",
         title="判不了，就说判不了",
         image="screenshots/13-决策工作台.png",
         bullets=["第 47 条弃权：批内只有 1 条资产，做不了跨资产查重",
                  "弃权单独列在【未决】里，不混进「通过」",
                  "工作台：左资产 / 中判定 / 右证据链",
                  "判定全部来自 std_verify，页面不算分"],
         script="这页是决策工作台。左边资产，中间判定，右边证据链。红色那条边，是实际命中的条款，连着它依据的标准。"
                "这里有个细节：第四十七条弃权了，因为这一批只有一条资产，跨资产查重需要至少两条。"
                "判不了就说判不了，不用「通过」糊过去。",
         action="指一下右侧红色那条边，再指黄色弃权框。"),
    dict(kind="flow",
         kicker="自进化闭环",
         title="它自己变好的方式：盲区 → 提案 → 四道闸门 → 归零",
         steps=["盲区", "候选判据", "四道闸门", "入库"],
         foot=["实测：盲区 2 → 0", "8 条候选被拒 —— 闸门在正常工作", "判据语言 8 类 → 12 类"],
         script="这是它自己变好的方式。系统发现盲区——就是明明违规、但现有判据看不见的资产。"
                "然后它提一条候选判据，交给四道闸门裁决：见过合规样本、零误报、能归因命中、不比原文松。过了才入库。"
                "实测里，盲区从两条变成零。也有八条候选被拒——这不是失败，是闸门在正常工作。",
         action="说到「八条被拒」时，语气放平，这是加分点。"),
    dict(kind="image_right",
         kicker="知识图谱",
         title="本体不是手画的，是机械导出的",
         image="kg/kg-schema.png",
         bullets=["8 类实体 / 8 种关系",
                  "145 节点 / 982 条边",
                  "每条边可回指：字段名 / 条款参数",
                  "条款→标准名回指 100%（41/41）"],
         script="本体不是我们手画的。八类实体、八种关系，从判据库和语料里机械导出，一百四十五个节点、九百八十二条边。"
                "每条边都能回指到字段或者条款参数。",
         action="指一下图中「条款→字段」那条线。"),
    dict(kind="image_right",
         kicker="多行业迁移",
         title="换行业只换一个包，判定逻辑不动",
         image="screenshots/15-决策工作台-政务领域.png",
         bullets=["制造业：38 判据 / 103 锚点",
                  "政务：23 判据 / 65 锚点",
                  "内核算法与 MCP 协议层 0 行改动",
                  "装配层约 15 行（如实说明）"],
         script="换行业要改多少？只换一个领域包。制造业是三十八条判据、一百零三条锚点；政务是二十三条、六十五条。"
                "内核算法和 MCP 协议层没动。右边是同一个工作台，切到政务的样子。",
         action="强调「0 行」和「15 行」的区别，别含糊。"),
    dict(kind="closing",
         kicker="边界（写在文档第一页）",
         title="三条我们没做到的事",
         bullets=["数据是虚构的：换成真实标准，所有数字必须重测",
                  "检索与图还没有完全打通",
                  "不宣称全自动，也不宣称替代专家"],
         foot="github.com/isheng-eqi/evostd · Apache-2.0",
         script="最后说边界。第一，数据是虚构的，标准是我们自己编的演示样例，换成真标准所有数字要重测。"
                "第二，检索和图还没完全打通。第三，我们不宣称全自动、也不宣称替代专家——三条不可编译的条款，系统明确弃权，留给人。"
                "代码已经开源，地址在下面。",
         action="三条逐条说，别赶。说完停一秒再结束。"),
]


NEW_SCRIPTS = [
    "大家好，我们是 WarrantedBelief 队，做的是赛题三。先说清楚：它不是让模型自己打分，判断依据来自标准原文。",
    "工厂和政府部门里，台账、点检表、图纸、规程堆了很多。不缺数据，缺的是「什么算对」。同一条台账，两个人能给出相反结论，谁也拿不出依据。",
    "我们的做法：把条款变成机器能跑的检查。比如检修周期不能超过九十天。再从原文里机械抽出「肯定合规、肯定违规」的样本，我们叫硬锚点。新规则想进库，得先在锚点上证明自己不误报、能抓到问题。最后拍板的是人。",
    "系统分三层：上面是 Nexent 智能体，负责选工具、组织语言；中间是我们自研的 MCP 服务，三十六个工具，判据和台账都在这；下面是知识库，只提供证据。模型能调用，改不了规则。",
    "这是平台上的配置。deepseek-flash，三十六个 MCP 工具，一个知识库检索，五个技能。每个工具都能点开调。",
    "问它 A 类设备的检修周期上限。它先去知识库检索，再查判据第四十五条。回答里两条依据都给：判据原文，语料出处，还带检索得分。检索找证据，判据判对错，这两件事不能混。",
    "再看一条决策：这条资产能不能入库。它依次调工具，结论是不能。依据第四十五条，A 类上限九十天，这条填了三百六十五天。证据字段原样贴出来。",
    "这是决策工作台，左资产、中判定、右证据链。注意这里：第四十七条弃权了，因为这一批只有一条资产，做不了跨资产查重。判不了就说判不了。",
    "它怎么自己变好？发现盲区，提候选判据，过四道闸门：见过合规样本、零误报、能归因、不比原文松。过了才入库。实测盲区从两条到零；也有八条候选被拒——那是闸门在工作。",
    "本体不是手画的。八类实体、八种关系，从判据和语料里机械导出，一百四十五个节点、九百八十二条边，每条都能回指。",
    "换行业只换一个包。制造业三十八条判据，政务二十三条。内核和协议层没动，装配层十几行。",
    "最后说边界。数据是虚构的，换真标准所有数字要重测；检索和图还没完全打通；我们不说全自动，也不说替代专家。代码开源在下面。",
]
for _s, _t in zip(SLIDES, NEW_SCRIPTS):
    _s["script"] = _t


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    total = len(SLIDES)
    for idx, s in enumerate(SLIDES, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
        dark = s["kind"] in ("cover", "closing")
        bg(slide, DARK if dark else WHITE)
        if dark:
            rect(slide, 0, H - Inches(0.12), W, Inches(0.12), ACCENT)

        if s["kind"] == "cover":
            rect(slide, Inches(0.9), Inches(1.2), Inches(0.9), Inches(0.07), ACCENT)
            textbox(slide, Inches(0.9), Inches(1.45), Inches(11), Inches(0.5),
                    [(s["kicker"], 18, RGBColor(0x9C, 0xC3, 0xFF), False)])
            textbox(slide, Inches(0.9), Inches(2.0), Inches(11), Inches(2.6),
                    [(line, 48, WHITE, True) for line in s["title"].split("\n")])
            textbox(slide, Inches(0.9), Inches(5.6), Inches(11), Inches(0.5),
                    [(s["sub"], 18, RGBColor(0xB8, 0xC2, 0xD0), False)])
            notes(slide, s["script"], s.get("action", ""))
            page_no(slide, idx, total, dark=True)
            continue

        # 顶部 kicker + 标题
        textbox(slide, Inches(0.8), Inches(0.5), Inches(11.8), Inches(0.4),
                [(s["kicker"], 14, ACCENT if not dark else RGBColor(0x9C, 0xC3, 0xFF), True)])
        title_color = WHITE if dark else INK
        textbox(slide, Inches(0.8), Inches(0.9), Inches(11.8), Inches(0.9),
                [(s["title"], 30, title_color, True)])
        rect(slide, Inches(0.8), Inches(1.72), Inches(1.1), Inches(0.05),
             ACCENT if not dark else RGBColor(0x54, 0xAE, 0xFF))

        has_image = s["kind"] in ("image_right", "diagram")
        text_w = Inches(5.6) if has_image else Inches(11.7)
        body_top = Inches(2.05) if not has_image else Inches(2.05)

        if s["kind"] == "statement":
            if s.get("big"):
                textbox(slide, Inches(0.8), Inches(2.0), Inches(11.7), Inches(1.0),
                        [(s["big"], 24, MUTED, False)])
            bullets = s.get("bullets") or []
            if s.get("chips"):
                x = Inches(0.8)
                for chip in s["chips"]:
                    box = rect(slide, x, Inches(3.4), Inches(2.2), Inches(0.7),
                               SOFT, RGBColor(0xD0, 0xD7, 0xDE), radius=True)
                    tf = box.text_frame
                    tf.margin_top = 0
                    tf.margin_bottom = 0
                    p = tf.paragraphs[0]
                    p.alignment = PP_ALIGN.CENTER
                    r = p.add_run()
                    r.text = chip
                    set_font(r, 16, INK, False)
                    x += Inches(2.5)
            start_y = Inches(3.05) if s.get("big") else Inches(2.15)
            for i, b in enumerate(bullets):
                textbox(slide, Inches(0.8), start_y + Inches(0.62) * i,
                        Inches(11.5), Inches(0.5), [("· " + b, 17, INK, False)])
        elif s["kind"] == "flow":
            x = Inches(0.8)
            y = Inches(2.3)
            for i, step in enumerate(s["steps"]):
                box = rect(slide, x, y, Inches(2.4), Inches(1.0), SOFT,
                           ACCENT if i in (0, 3) else RGBColor(0xD0, 0xD7, 0xDE), radius=True)
                tf = box.text_frame
                tf.margin_top = 0
                tf.margin_bottom = 0
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                r = p.add_run()
                r.text = step
                set_font(r, 20 if i != 2 else 18, INK, True)
                if i < len(s["steps"]) - 1:
                    textbox(slide, x + Inches(2.45), y + Inches(0.18), Inches(0.5), Inches(0.6),
                            [("→", 28, ACCENT, True)])
                x += Inches(2.95)
            for i, f in enumerate(s.get("foot", [])):
                textbox(slide, Inches(0.8), Inches(4.1) + Inches(0.55) * i,
                        Inches(11.5), Inches(0.5), [("· " + f, 17, INK if i == 0 else MUTED, False)])
        elif s["kind"] == "image_bottom":
            for i, b in enumerate(s["bullets"]):
                col, row = i % 2, i // 2
                textbox(slide, Inches(0.8) + Inches(5.9) * col,
                        Inches(1.95) + Inches(0.52) * row,
                        Inches(5.6), Inches(0.5), [("· " + b, 15, INK, False)])
        elif s["kind"] != "closing":
            if s.get("bullets"):
                for i, b in enumerate(s["bullets"]):
                    textbox(slide, Inches(0.8), body_top + Inches(0.68) * i,
                            text_w, Inches(0.62), [("· " + b, 17, INK, False)])

        if s["kind"] == "diagram":
            layers = [
                ("Nexent 智能体", "deepseek-flash · 只编排、不判定", RGBColor(0xE8, 0xF0, 0xFE), ACCENT),
                ("evostd MCP Server", "36 工具 · 判据 / 锚点 / 台账", RGBColor(0xE8, 0xF5, 0xEC), GREEN),
                ("知识库 mfg_assets", "10 文档 / 51 分块 · 只提供证据", RGBColor(0xF3, 0xEE, 0xFB), RGBColor(0x7C, 0x3A, 0xED)),
            ]
            y = Inches(2.05)
            for i, (name, desc, fill, line) in enumerate(layers):
                box = rect(slide, Inches(6.9), y, Inches(5.7), Inches(1.05), fill, line, radius=True)
                tf = box.text_frame
                tf.word_wrap = True
                tf.margin_top = 0
                tf.margin_bottom = 0
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                r = p.add_run()
                r.text = name
                set_font(r, 19, INK, True)
                p2 = tf.add_paragraph()
                p2.alignment = PP_ALIGN.CENTER
                r2 = p2.add_run()
                r2.text = desc
                set_font(r2, 13, MUTED, False)
                if i < len(layers) - 1:
                    textbox(slide, Inches(9.35), y + Inches(1.08), Inches(1.4), Inches(0.4),
                            [("↓ 调用" if i == 0 else "↓ 检索", 13, MUTED, False)],
                            align=PP_ALIGN.CENTER)
                y += Inches(1.5)
            textbox(slide, Inches(6.9), Inches(6.55), Inches(5.7), Inches(0.4),
                    [("判定只发生在中间层：模型看得到结论，改不了规则", 12, MUTED, False)])

        if s["kind"] == "image_right":
            image_cropped(slide, s["image"], Inches(6.75), Inches(1.95),
                          Inches(5.85), Inches(4.9), s.get("crop"))
        elif s["kind"] == "image_bottom":
            image_cropped(slide, s["image"], Inches(0.8), Inches(3.35),
                          Inches(11.7), Inches(3.7), s.get("crop"))

        if s["kind"] == "closing":
            for i, b in enumerate(s["bullets"]):
                textbox(slide, Inches(0.8), Inches(2.2) + Inches(0.75) * i,
                        Inches(11.5), Inches(0.6), [("· " + b, 20, WHITE, False)])
            textbox(slide, Inches(0.8), Inches(5.2), Inches(11.5), Inches(0.5),
                    [(s["foot"], 16, RGBColor(0x9C, 0xC3, 0xFF), False)])

        notes(slide, s["script"], s.get("action", ""))
        page_no(slide, idx, total, dark=dark)

    prs.save(OUT)
    print(f"[ppt] 已生成 {OUT}（{total} 页）")


if __name__ == "__main__":
    build()
